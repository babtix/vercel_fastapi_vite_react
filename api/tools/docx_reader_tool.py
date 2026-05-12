"""Tool for extracting text from Word documents (.docx).

Uses the python-docx library to parse paragraphs and tables.
Extraction runs inside a threadpool to avoid blocking the async event loop.
"""

import io
import docx
import logging
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, status
from fastapi.concurrency import run_in_threadpool
from pydantic import BaseModel
from typing import Dict
from dependencies import get_current_user

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/tools/docx", tags=["Tools"])


class DocxExtractResponse(BaseModel):
    """Schema returned after extracting text from a DOCX file.

    Attributes:
        status: Result status — typically "success".
        filename: Original name of the uploaded file.
        paragraphs_count: Number of paragraphs found in the document.
        tables_count: Number of tables found in the document.
        text: Concatenated extracted text from paragraphs and table rows.
    """

    status: str
    filename: str
    paragraphs_count: int
    tables_count: int
    text: str


def _extract_docx_text_sync(file_stream) -> dict:
    """Synchronously extract text from a DOCX file stream.

    Iterates over all paragraphs and table rows, collecting non-empty text.

    Args:
        file_stream: A file-like object containing DOCX data.

    Returns:
        Dictionary with paragraphs_count, tables_count, and the combined text.
    """
    doc = docx.Document(file_stream)
    extracted_text = []

    # Collect text from all non-empty paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            extracted_text.append(para.text.strip())

    # Collect text from all table cells, joining each row with " | "
    for table in doc.tables:
        for row in table.rows:
            row_data = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if row_data:
                extracted_text.append(" | ".join(row_data))

    full_text = "\n\n".join(extracted_text)

    return {
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
        "text": full_text,
    }


@router.post(
    "/extract-text",
    response_model=DocxExtractResponse,
    status_code=status.HTTP_200_OK,
)
async def extract_docx_text(
    file: UploadFile = File(
        ..., description="Word file (.docx) to read"
    ),
    current_user: dict = Depends(get_current_user),
):
    """Extract plain text from an uploaded DOCX file.

    Args:
        file: Uploaded DOCX file.
        current_user: Injected authenticated user dependency.

    Returns:
        A DocxExtractResponse with paragraph/table counts and extracted text.

    Raises:
        HTTPException: 400 if the file is not a DOCX.
        HTTPException: 500 if extraction fails internally.
    """
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Le fichier doit avoir l'extension .docx",
        )

    try:
        # Read the upload into memory and wrap it for python-docx
        content = await file.read()
        file_stream = io.BytesIO(content)

        # Offload CPU-bound DOCX parsing to a background thread
        result = await run_in_threadpool(
            _extract_docx_text_sync, file_stream
        )

        return DocxExtractResponse(
            status="success",
            filename=file.filename,
            **result,
        )
    except Exception as exc:
        logger.error("DOCX extraction failed: %s", exc)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Échec de lecture du fichier DOCX: {str(exc)}",
        )
