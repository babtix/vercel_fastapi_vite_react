"""RAG (Retrieval-Augmented Generation) Service.

Handles the full document lifecycle for agent knowledge bases:
- Text extraction from PDF, DOCX, and plain-text files
- Chunking using RecursiveCharacterTextSplitter
- Embedding generation via Ollama
- Vector storage and semantic search with ChromaDB
"""

import os
import io
import hashlib
import logging
import chromadb
import ollama as ollama_client
import pypdf
import docx
from typing import List, Dict, Optional
from langchain_text_splitters import RecursiveCharacterTextSplitter
from core.settings import settings

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# ChromaDB persistent storage setup
# ---------------------------------------------------------------------------
CHROMA_DIR = os.path.join(
    os.path.dirname(os.path.dirname(__file__)), "chroma_data"
)
os.makedirs(CHROMA_DIR, exist_ok=True)
chroma = chromadb.PersistentClient(path=CHROMA_DIR)

# ---------------------------------------------------------------------------
# Ollama client initialization (used for embedding generation)
# ---------------------------------------------------------------------------
_ollama_client = ollama_client.Client(
    host=settings.OLLAMA_URL, timeout=settings.OLLAMA_TIMEOUT
)

# ---------------------------------------------------------------------------
# Text splitter configuration
# Splits documents into overlapping chunks to preserve context across boundaries.
# ---------------------------------------------------------------------------
splitter = RecursiveCharacterTextSplitter(
    chunk_size=settings.RAG_CHUNK_SIZE,
    chunk_overlap=settings.RAG_CHUNK_OVERLAP,
    separators=["\n\n", "\n", ". ", " ", ""],
)


def _collection_name(agent_id: str) -> str:
    """Generate a deterministic ChromaDB collection name for an agent."""
    return f"agent_{agent_id}"


def _get_or_create_collection(agent_id: str):
    """Retrieve an existing ChromaDB collection or create one with cosine similarity."""
    return chroma.get_or_create_collection(
        name=_collection_name(agent_id),
        metadata={"hnsw:space": "cosine"},
    )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF byte stream using pypdf."""
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return "\n\n".join(pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX byte stream, including paragraphs and table rows."""
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_data = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if row_data:
                parts.append(" | ".join(row_data))
    return "\n\n".join(parts)


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decode raw bytes to a UTF-8 string, replacing invalid characters."""
    return file_bytes.decode("utf-8", errors="replace")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Route text extraction to the appropriate parser based on file extension.

    Args:
        filename: Original filename (used to infer type).
        file_bytes: Raw file contents.

    Returns:
        Extracted plain text.

    Raises:
        ValueError: If the file extension is not supported.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        return extract_text_from_docx(file_bytes)
    elif ext in ("txt", "md", "csv", "json", "xml", "html"):
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")


def _embed_texts(texts: List[str]) -> List[List[float]]:
    """Generate embedding vectors for a list of texts using Ollama.

    Processes texts in batches of 32 to avoid overloading the embedding endpoint.

    Args:
        texts: List of text strings to embed.

    Returns:
        List of embedding vectors (each vector is a list of floats).
    """
    all_embeddings = []
    batch_size = 32

    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        response = _ollama_client.embed(
            model=settings.RAG_EMBEDDING_MODEL, input=batch
        )
        all_embeddings.extend(response["embeddings"])

    return all_embeddings


def _embed_query(query: str) -> List[float]:
    """Embed a single query string into a vector."""
    return _embed_texts([query])[0]


def ingest_document(
    agent_id: str, filename: str, file_bytes: bytes
) -> Dict:
    """Run the full document ingestion pipeline for an agent.

    Pipeline steps:
      1. Extract raw text from the uploaded file.
      2. Split text into overlapping chunks.
      3. Compute an MD5 hash to identify the document.
      4. Generate embedding vectors for each chunk.
      5. Store chunks + embeddings + metadata in ChromaDB.

    Args:
        agent_id: MongoDB ObjectId of the target agent.
        filename: Original filename (stored in metadata).
        file_bytes: Raw file contents.

    Returns:
        Dictionary with ingestion metadata (filename, doc_hash, total_chars, chunks_count).

    Raises:
        ValueError: If the document is empty or yields no chunks after splitting.
    """
    text = extract_text(filename, file_bytes)
    if not text.strip():
        logger.warning("Empty or unreadable document uploaded: %s", filename)
        raise ValueError("Le document est vide ou illisible.")

    # Split the extracted text into semantically coherent chunks
    chunks = splitter.split_text(text)
    if not chunks:
        raise ValueError(
            "Aucun contenu exploitable après découpage du document."
        )

    # Use the first 8 characters of the MD5 digest as a short document identifier
    doc_hash = hashlib.md5(file_bytes).hexdigest()[:8]
    ids = [f"{doc_hash}_{i}" for i in range(len(chunks))]

    # Generate embeddings for all chunks
    embeddings = _embed_texts(chunks)

    # Persist to the agent's ChromaDB collection
    collection = _get_or_create_collection(agent_id)
    metadatas = [
        {"filename": filename, "doc_hash": doc_hash, "chunk_index": i}
        for i in range(len(chunks))
    ]

    collection.add(
        ids=ids,
        embeddings=embeddings,
        documents=chunks,
        metadatas=metadatas,
    )

    return {
        "filename": filename,
        "doc_hash": doc_hash,
        "total_chars": len(text),
        "chunks_count": len(chunks),
    }


def retrieve_context(
    agent_id: str, query: str, top_k: int = None
) -> List[Dict]:
    """Perform semantic search to retrieve relevant document chunks for a query.

    Args:
        agent_id: MongoDB ObjectId of the agent whose knowledge base to search.
        query: Natural language query string.
        top_k: Maximum number of chunks to retrieve (default from settings).

    Returns:
        List of chunk dicts containing text, filename, and similarity score.
    """
    if top_k is None:
        top_k = settings.RAG_TOP_K

    collection = _get_or_create_collection(agent_id)
    if collection.count() == 0:
        return []

    # Embed the query so it can be compared against stored document vectors
    query_embedding = _embed_query(query)

    # Query ChromaDB for the nearest neighbors in embedding space
    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=min(top_k, collection.count()),
        include=["documents", "metadatas", "distances"],
    )

    # Reconstruct result objects with cosine-similarity-like scores
    context_chunks = []
    for i in range(len(results["ids"][0])):
        context_chunks.append(
            {
                "text": results["documents"][0][i],
                "filename": results["metadatas"][0][i].get(
                    "filename", "unknown"
                ),
                # Convert distance to a similarity score (1 - distance)
                "score": 1 - results["distances"][0][i],
            }
        )

    # Filter out low-quality matches below the configured minimum score threshold
    context_chunks = [
        c for c in context_chunks if c["score"] >= settings.RAG_MIN_SCORE
    ]
    return context_chunks


def build_rag_prompt(context_chunks: List[Dict]) -> str:
    """Build a formatted context block from retrieved chunks for injection into the system prompt.

    Args:
        context_chunks: List of chunk dicts returned by retrieve_context.

    Returns:
        A multi-line string containing the formatted RAG context block,
        or an empty string if no chunks are provided.
    """
    if not context_chunks:
        return ""

    context_parts = []
    for i, chunk in enumerate(context_chunks, 1):
        context_parts.append(
            f"[Source: {chunk['filename']}]\n{chunk['text']}"
        )

    context_block = "\n\n---\n\n".join(context_parts)

    return (
        "\n\n"
        "══════════════════════════════════════════\n"
        "CONTEXTE DOCUMENTAIRE (Base de connaissances)\n"
        "══════════════════════════════════════════\n"
        "Utilise les informations ci-dessous pour répondre à la question "
        "de l'utilisateur. Si la réponse n'est pas dans le contexte, "
        "dis-le clairement.\n\n"
        f"{context_block}\n"
        "══════════════════════════════════════════\n"
    )


def list_documents(agent_id: str) -> List[Dict]:
    """List all unique documents indexed for a given agent.

    Args:
        agent_id: MongoDB ObjectId of the agent.

    Returns:
        List of dicts containing filename, doc_hash, and chunk count per document.
    """
    collection = _get_or_create_collection(agent_id)
    if collection.count() == 0:
        return []

    all_data = collection.get(include=["metadatas"])
    doc_map = {}
    # Aggregate chunks by doc_hash to produce per-document summaries
    for meta in all_data["metadatas"]:
        key = meta.get("doc_hash", "unknown")
        if key not in doc_map:
            doc_map[key] = {
                "filename": meta.get("filename", "unknown"),
                "doc_hash": key,
                "chunks_count": 0,
            }
        doc_map[key]["chunks_count"] += 1

    return list(doc_map.values())


def delete_document(agent_id: str, doc_hash: str) -> int:
    """Delete all chunks belonging to a specific document from an agent's collection.

    Args:
        agent_id: MongoDB ObjectId of the agent.
        doc_hash: Short hash identifying the document to remove.

    Returns:
        Number of chunks that were deleted.
    """
    collection = _get_or_create_collection(agent_id)
    if collection.count() == 0:
        return 0

    # Find all chunk IDs matching the doc_hash before deletion
    matching = collection.get(
        where={"doc_hash": doc_hash}, include=["metadatas"]
    )
    count = len(matching["ids"])

    if count > 0:
        collection.delete(where={"doc_hash": doc_hash})

    return count


def delete_all_documents(agent_id: str) -> int:
    """Drop the entire ChromaDB collection for an agent.

    Args:
        agent_id: MongoDB ObjectId of the agent.

    Returns:
        Number of chunks that existed before deletion.
    """
    collection_name = _collection_name(agent_id)
    try:
        collection = chroma.get_collection(collection_name)
        count = collection.count()
        chroma.delete_collection(collection_name)
        return count
    except Exception as exc:
        logger.warning("Could not delete collection '%s': %s", collection_name, exc)
        return 0


def get_agent_doc_count(agent_id: str) -> int:
    """Return the total number of indexed chunks for an agent.

    Args:
        agent_id: MongoDB ObjectId of the agent.

    Returns:
        Chunk count, or 0 if the collection does not exist or is unreachable.
    """
    try:
        collection = _get_or_create_collection(agent_id)
        return collection.count()
    except Exception as exc:
        logger.warning("Could not count docs for agent '%s': %s", agent_id, exc)
        return 0

    matching = collection.get(
        where={"doc_hash": doc_hash}, include=["metadatas"]
    )
    count = len(matching["ids"])

    if count > 0:
        collection.delete(where={"doc_hash": doc_hash})

    return count


def delete_all_documents(agent_id: str) -> int:
    # Delete the entire collection for an agent
    collection_name = _collection_name(agent_id)
    try:
        collection = chroma.get_collection(collection_name)
        count = collection.count()
        chroma.delete_collection(collection_name)
        return count
    except Exception as exc:
        logger.warning("Could not delete collection '%s': %s", collection_name, exc)
        return 0


def get_agent_doc_count(agent_id: str) -> int:
    # Get total chunk count for an agent
    try:
        collection = _get_or_create_collection(agent_id)
        return collection.count()
    except Exception as exc:
        logger.warning("Could not count docs for agent '%s': %s", agent_id, exc)
        return 0
